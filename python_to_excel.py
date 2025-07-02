import json
import os
import sys
from pathlib import Path

import openpyxl
import pandas as pd


def python_to_excel(project_dir: str = "output", excel_path: str = "reconstructed.xlsx") -> None:
    data_dir = Path(project_dir) / "data"
    formula_dir = Path(project_dir) / "formulas"
    wb = openpyxl.Workbook()
    if wb.worksheets:
        wb.remove(wb.active)

    sheets = {}

    for data_file in data_dir.glob("*_data_block*.json"):
        name_part = data_file.stem
        sheet_name = name_part.split("_data_block")[0]
        ws = sheets.get(sheet_name)
        if ws is None:
            ws = wb.create_sheet(title=sheet_name)
            sheets[sheet_name] = ws
        df = pd.read_json(data_file, orient="split")
        for r in df.index:
            for c in df.columns:
                ws[f"{c}{r}"] = df.at[r, c]

    for formula_file in formula_dir.glob("*_block*.json"):
        name_part = formula_file.stem
        sheet_name = name_part.split("_block")[0]
        ws = sheets.get(sheet_name)
        if ws is None:
            ws = wb.create_sheet(title=sheet_name)
            sheets[sheet_name] = ws
        with open(formula_file, "r", encoding="utf-8") as f:
            formulas = json.load(f)
        for cell, formula in formulas.items():
            ws[cell] = formula
    wb.save(excel_path)


if __name__ == "__main__":
    project = sys.argv[1] if len(sys.argv) > 1 else "output"
    out_excel = sys.argv[2] if len(sys.argv) > 2 else "reconstructed.xlsx"
    python_to_excel(project, out_excel)
